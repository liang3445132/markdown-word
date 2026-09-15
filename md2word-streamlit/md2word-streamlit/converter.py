from __future__ import annotations

import io
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree


BODY_FONT = "宋体"
BODY_FONT_ASCII = "SimSun"
BODY_SIZE_PT = 10.5
BODY_STYLE_NAME = "正文"
HEADING_FONT = "黑体"
HEADING_FONT_ASCII = "SimHei"
MAX_MARKDOWN_CHARS = 300_000


class ConversionError(RuntimeError):
    """A user-facing conversion failure."""


def _set_style_font(style, east_asia: str, ascii_font: str, size_pt: float, bold: bool | None = None) -> None:
    style.font.name = ascii_font
    style.font.size = Pt(size_pt)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.italic = False
    style.font.underline = False
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), ascii_font)


def _set_run_font(run, east_asia: str = BODY_FONT, ascii_font: str = BODY_FONT_ASCII, size_pt: float = BODY_SIZE_PT) -> None:
    run.font.name = ascii_font
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), ascii_font)


def _add_default_run_properties(document: Document) -> None:
    styles_root = document.styles.element
    doc_defaults = styles_root.find(qn("w:docDefaults"))
    if doc_defaults is None:
        doc_defaults = OxmlElement("w:docDefaults")
        styles_root.insert(0, doc_defaults)
    rpr_default = doc_defaults.find(qn("w:rPrDefault"))
    if rpr_default is None:
        rpr_default = OxmlElement("w:rPrDefault")
        doc_defaults.append(rpr_default)
    rpr = rpr_default.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_default.append(rpr)

    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr, value in (("ascii", BODY_FONT_ASCII), ("hAnsi", BODY_FONT_ASCII), ("eastAsia", BODY_FONT), ("cs", BODY_FONT_ASCII)):
        rfonts.set(qn(f"w:{attr}"), value)

    for tag in ("w:sz", "w:szCs"):
        element = rpr.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            rpr.append(element)
        element.set(qn("w:val"), "21")


def _ensure_body_style(document: Document):
    if BODY_STYLE_NAME in document.styles:
        style = document.styles[BODY_STYLE_NAME]
    else:
        style = document.styles.add_style(BODY_STYLE_NAME, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["Normal"]
    _set_style_font(style, BODY_FONT, BODY_FONT_ASCII, BODY_SIZE_PT)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    style.paragraph_format.space_after = Pt(4)
    return style


def build_reference_docx(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.6)

    _add_default_run_properties(document)
    _ensure_body_style(document)

    body_styles = {
        "Normal",
        "Body Text",
        "Body Text 2",
        "Body Text 3",
        "First Paragraph",
        "Compact",
        "Table",
        "Caption",
        "Block Text",
    }
    for style in document.styles:
        if style.type in {WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER} and style.name in body_styles:
            _set_style_font(style, BODY_FONT, BODY_FONT_ASCII, BODY_SIZE_PT)
            if style.type == WD_STYLE_TYPE.PARAGRAPH:
                style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
                style.paragraph_format.space_after = Pt(4)

    heading_sizes = {"Title": 20, "Subtitle": 12, "Heading 1": 16, "Heading 2": 14, "Heading 3": 12, "Heading 4": 10.5}
    for name, size in heading_sizes.items():
        if name in document.styles:
            _set_style_font(document.styles[name], HEADING_FONT, HEADING_FONT_ASCII, size, bold=True)

    for name in ("Source Code", "Verbatim Char"):
        if name in document.styles:
            _set_style_font(document.styles[name], "等线", "Consolas", 9)

    document.core_properties.title = "墨转 Markdown 文档"
    document.core_properties.subject = "Markdown 转 Word"
    document.core_properties.author = "墨转"
    document.save(path)


def _normalize_math_delimiters(markdown: str) -> str:
    markdown = re.sub(
        r"(?ms)^```(?:latex|tex|math)\s*\n([\s\S]*?)\n```[ \t]*$",
        lambda match: f"$$\n{match.group(1).strip()}\n$$",
        markdown,
    )
    markdown = re.sub(
        r"(?ms)^~~~(?:latex|tex|math)\s*\n([\s\S]*?)\n~~~[ \t]*$",
        lambda match: f"$$\n{match.group(1).strip()}\n$$",
        markdown,
    )
    chunks = re.split(r"(```[\s\S]*?```|~~~[\s\S]*?~~~|`[^`\n]*`)", markdown)
    for index in range(0, len(chunks), 2):
        chunks[index] = (
            chunks[index]
            .replace(r"\\[", "$$")
            .replace(r"\\]", "$$")
            .replace(r"\\(", "$")
            .replace(r"\\)", "$")
            .replace(r"\[", "$$")
            .replace(r"\]", "$$")
            .replace(r"\(", "$")
            .replace(r"\)", "$")
        )
    normalized = "".join(chunks)

    math_parts = re.split(r"(\$\$[\s\S]*?\$\$|(?<!\$)\$(?!\$).*?(?<!\$)\$(?!\$))", normalized)
    operator_wrappers = re.compile(r"\\mathrm\s*\{\s*([=+\-(),;:])\s*\}")
    for index in range(1, len(math_parts), 2):
        # Collapse doubled command slashes from copied/escaped Markdown while
        # preserving matrix row separators, which are not followed by letters.
        math_parts[index] = re.sub(r"\\\\(?=[A-Za-z]+(?:\b|\s*\{))", r"\\", math_parts[index])
        math_parts[index] = operator_wrappers.sub(r"\1", math_parts[index])
    return "".join(math_parts)


def _count_markdown_math(markdown: str) -> int:
    without_code = re.sub(r"```[\s\S]*?```|~~~[\s\S]*?~~~|`[^`\n]*`", "", markdown)
    display = re.findall(r"(?<!\\)\$\$([\s\S]+?)(?<!\\)\$\$", without_code)
    without_display = re.sub(r"(?<!\\)\$\$[\s\S]+?(?<!\\)\$\$", "", without_code)
    inline = re.findall(r"(?<![\\$])\$(?!\$)(.+?)(?<!\\)\$(?!\$)", without_display)
    return len(display) + len(inline)


def _pandoc_path() -> str:
    system_pandoc = shutil.which("pandoc")
    if system_pandoc:
        return system_pandoc
    try:
        import pypandoc

        return pypandoc.get_pandoc_path()
    except (ImportError, OSError) as exc:
        raise ConversionError(
            "转换引擎尚未就绪。请安装 pypandoc-binary，或在系统中安装 Pandoc 后重试。"
        ) from exc


def _iter_table_paragraphs(document: Document):
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                yield from cell.paragraphs
                for nested_table in cell.tables:
                    for nested_row in nested_table.rows:
                        for nested_cell in nested_row.cells:
                            yield from nested_cell.paragraphs


def _enforce_body_formatting(docx_path: Path, title: str) -> None:
    document = Document(docx_path)
    _add_default_run_properties(document)
    body_style = _ensure_body_style(document)
    document.core_properties.title = title
    document.core_properties.subject = "Markdown 转 Word"
    document.core_properties.author = "墨转"

    paragraphs = list(document.paragraphs) + list(_iter_table_paragraphs(document))
    heading_sizes = {"Title": 20, "Subtitle": 12, "Heading 1": 16, "Heading 2": 14, "Heading 3": 12, "Heading 4": 10.5}
    for paragraph in paragraphs:
        style_name = paragraph.style.name if paragraph.style else "Normal"
        if style_name in heading_sizes:
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                _set_run_font(run, east_asia=HEADING_FONT, ascii_font=HEADING_FONT_ASCII, size_pt=heading_sizes[style_name])
                run.font.bold = True
                run.font.italic = False
                run.font.underline = False
            continue
        if style_name not in {"Source Code", "Verbatim Char"}:
            paragraph.style = body_style
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        for run in paragraph.runs:
            if style_name in {"Source Code", "Verbatim Char"}:
                _set_run_font(run, east_asia="等线", ascii_font="Consolas", size_pt=9)
            else:
                _set_run_font(run)
    document.save(docx_path)


def _repair_strict_ooxml(docx_path: Path) -> None:
    """Repair Pandoc constructs that Word accepts but strict ISO schemas reject."""
    math_ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    word_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    with zipfile.ZipFile(docx_path, "r") as source:
        entries = [(item, source.read(item.filename)) for item in source.infolist()]

    repaired: list[tuple[zipfile.ZipInfo, bytes]] = []
    for item, data in entries:
        if item.filename == "word/document.xml":
            root = etree.fromstring(data)
            for properties in root.xpath(".//m:mcPr", namespaces={"m": math_ns}):
                count = properties.find(f"{{{math_ns}}}count")
                justification = properties.find(f"{{{math_ns}}}mcJc")
                if count is not None and justification is not None and properties.index(count) > properties.index(justification):
                    properties.remove(count)
                    properties.insert(properties.index(justification), count)
            data = etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)
        elif item.filename == "word/settings.xml":
            root = etree.fromstring(data)
            zoom = root.find(f"{{{word_ns}}}zoom")
            if zoom is not None and f"{{{word_ns}}}percent" not in zoom.attrib:
                zoom.set(f"{{{word_ns}}}percent", "100")
            data = etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)
        repaired.append((item, data))

    replacement = docx_path.with_suffix(".repaired.docx")
    with zipfile.ZipFile(replacement, "w") as target:
        for item, data in repaired:
            target.writestr(item, data)
    replacement.replace(docx_path)


def inspect_docx(docx_bytes: bytes) -> dict[str, int | bool]:
    try:
        with zipfile.ZipFile(io.BytesIO(docx_bytes)) as archive:
            bad_file = archive.testzip()
            if bad_file:
                raise ConversionError(f"生成的 Word 压缩包已损坏：{bad_file}")
            document_xml = archive.read("word/document.xml")
            styles_xml = archive.read("word/styles.xml")
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ConversionError("生成结果不是有效的 DOCX 文件。") from exc

    equation_count = len(re.findall(br"<m:oMath(?:\s|>)", document_xml))
    document_root = etree.fromstring(document_xml)
    text_nodes = document_root.xpath(
        ".//w:t/text() | .//m:t/text()",
        namespaces={
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
            "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
        },
    )
    visible_text = "\n".join(text_nodes)
    linear_latex_count = len(
        re.findall(
            r"\\(?:frac|dfrac|tfrac|sqrt|sum|prod|int|begin|end|mathrm|mathbf|text|hat|bar|vec)\b",
            visible_text,
        )
    )
    has_simsun = BODY_FONT.encode("utf-8") in styles_xml or BODY_FONT_ASCII.encode("ascii") in styles_xml
    has_five_size = b'w:val="21"' in styles_xml
    return {
        "equation_count": equation_count,
        "linear_latex_count": linear_latex_count,
        "has_simsun": has_simsun,
        "has_five_size": has_five_size,
    }


def convert_markdown_to_docx(markdown: str, title: str = "Markdown 文档") -> bytes:
    if not markdown.strip():
        raise ConversionError("请先输入 Markdown 内容。")
    if len(markdown) > MAX_MARKDOWN_CHARS:
        raise ConversionError(f"内容过长，请控制在 {MAX_MARKDOWN_CHARS:,} 个字符以内。")
    if "\x00" in markdown:
        raise ConversionError("内容包含不支持的空字符。")

    normalized = _normalize_math_delimiters(markdown)
    expected_math = _count_markdown_math(normalized)

    with tempfile.TemporaryDirectory(prefix="md2word-") as temp_dir:
        temp_path = Path(temp_dir)
        reference_path = temp_path / "reference.docx"
        output_path = temp_path / "document.docx"
        build_reference_docx(reference_path)

        command = [
            _pandoc_path(),
            "--from=markdown+tex_math_dollars+pipe_tables+task_lists+strikeout+footnotes+fenced_code_blocks",
            "--to=docx",
            f"--reference-doc={reference_path}",
            "--wrap=none",
            "--metadata=lang:zh-CN",
            f"--output={output_path}",
        ]
        try:
            completed = subprocess.run(
                command,
                input=normalized,
                text=True,
                encoding="utf-8",
                errors="strict",
                capture_output=True,
                timeout=45,
                check=False,
                cwd=temp_path,
            )
        except subprocess.TimeoutExpired as exc:
            raise ConversionError("内容转换超时，请缩短文稿或减少复杂公式后重试。") from exc
        except OSError as exc:
            raise ConversionError("无法启动 Word 转换引擎，请检查 Pandoc 是否可用。") from exc

        if completed.returncode != 0 or not output_path.exists():
            detail = (completed.stderr or completed.stdout or "未知错误").strip().splitlines()[-1]
            raise ConversionError(f"转换失败：{detail[:300]}")

        _enforce_body_formatting(output_path, title)
        _repair_strict_ooxml(output_path)
        result = output_path.read_bytes()
        report = inspect_docx(result)

    if not report["has_simsun"] or not report["has_five_size"]:
        raise ConversionError("字体校验未通过：未能写入宋体五号样式。")
    if expected_math and report["equation_count"] == 0:
        raise ConversionError("公式校验未通过：公式没有转换为 Word 原生可编辑对象，请检查 LaTeX 写法。")
    if report["linear_latex_count"]:
        raise ConversionError(
            "公式校验未通过：仍有公式处于 Word 线性模式，未生成专业格式。"
            "请使用 $...$、$$...$$ 或 latex 代码块包裹公式。"
        )
    return result
